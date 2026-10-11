"""JSON API used by the two dashboards."""
import csv
import io
import re
from collections import Counter

from flask import Blueprint, current_app, g, jsonify, request

from . import auth, stats
from .queries import (ApiError, clean_ip_phone, clean_note, clean_reason,
                      decorate, filtered_records, json_body, paginate, parse_day, parse_range,
                      reasons_by_key)
from .security import generate_password, hash_password, password_problem
from .storage import (STATUSES, STATUS_LABELS, StorageError, iso_now, new_record_id, norm_uid,
                      parse_record_id, valid_uid)

bp = Blueprint("api", __name__)

ROLES = ("user", "admin")


def public_user(u):
    return {"user_id": u["user_id"], "name": u["name"], "department": u.get("department", ""),
            "role": u["role"]}


def admin_user_row(u):
    return {**public_user(u), "active": u.get("active", True),
            "must_change_password": bool(u.get("must_change_password")),
            "last_login": u.get("last_login"), "created_at": u.get("created_at")}


def edit_window_open(day):
    """May a non-admin still change a record from ``day``?"""
    limit = current_app.config["USER_EDIT_DAYS"]
    if limit is None:
        return True
    from datetime import date
    return (date.fromisoformat(auth.today()) - date.fromisoformat(day)).days <= int(limit)


# ====================================================================== common
@bp.get("/me")
@auth.login_required
def me():
    return jsonify(
        user=public_user(g.user),
        reasons=auth.store().reasons(),
        status_labels=STATUS_LABELS,
        csrf=auth.csrf_token(),
        today=auth.today(),
        server_time=iso_now(auth.now()),
        user_edit_days=current_app.config["USER_EDIT_DAYS"],
    )


# =================================================================== my records
def day_payload(day, uid):
    users = auth.store().users()
    reasons = reasons_by_key()
    recs = auth.store().day_records(day, uid)
    rows = [decorate(r, users, reasons) for r in recs]
    rows.sort(key=lambda r: r["time"], reverse=True)
    editable = edit_window_open(day)
    for row in rows:
        row["editable"] = editable
    summary = stats.summarize(recs, reasons)
    return {
        "date": day, "server_time": iso_now(auth.now()), "records": rows,
        "stats": {**{k: summary[k] for k in ("total", "first_call", "last_call", "by_reason", *STATUSES)},
                  "hourly": stats.buckets(recs, reasons, "hour")},
    }


@bp.get("/records")
@auth.login_required
def my_records():
    day = parse_day(request.args.get("date"), auth.today())
    if day > auth.today():
        raise ApiError("That date is in the future.")
    return jsonify(day_payload(day, g.user["user_id"]))


@bp.post("/records")
@auth.login_required
def create_record():
    if g.user["role"] != "user":
        raise ApiError("Administrator accounts do not log calls. Use an employee account.", 403)
    body = json_body()
    ip_phone = clean_ip_phone(body.get("ip_phone"))
    reason = clean_reason(body.get("reason"))
    note = clean_note(body.get("note"))
    stamp = auth.now().replace(microsecond=0)
    day = stamp.date().isoformat()
    rec = {
        "id": new_record_id(day, g.user["user_id"]), "user_id": g.user["user_id"],
        "date": day, "time": stamp.time().isoformat(), "ip_phone": ip_phone,
        "reason": reason, "note": note, "created_at": iso_now(stamp),
        "updated_at": None, "updated_by": None,
    }
    auth.store().add_record(rec)
    auth.audit("record_created", target=rec["id"],
               details={"ip_phone": ip_phone, "reason": reason, "note": note})
    return jsonify(record=decorate(rec, auth.store().users(), reasons_by_key())), 201


def _owned_record_ref(rid):
    """(day, owner) of ``rid`` if the signed-in user may touch it, else 404 / 403."""
    parsed = parse_record_id(rid)
    if not parsed:
        raise ApiError("Record not found.", 404)
    day, owner = parsed
    is_admin = g.user["role"] == "admin"
    if owner != g.user["user_id"] and not is_admin:
        raise ApiError("Record not found.", 404)
    if not is_admin and not edit_window_open(day):
        raise ApiError(
            f"Records older than {current_app.config['USER_EDIT_DAYS']} days are locked. "
            "Ask your administrator.", 403)
    return day, owner


@bp.put("/records/<rid>")
@auth.login_required
def update_record(rid):
    _, owner = _owned_record_ref(rid)
    body = json_body()
    changes = {}
    if "ip_phone" in body:
        changes["ip_phone"] = clean_ip_phone(body["ip_phone"])
    if "reason" in body:
        changes["reason"] = clean_reason(body["reason"])
    if "note" in body:
        changes["note"] = clean_note(body["note"])
    if not changes:
        raise ApiError("Nothing to change.")
    current = auth.store().get_record(rid)
    if not current:
        raise ApiError("Record not found.", 404)
    changes = {k: v for k, v in changes.items() if current.get(k) != v}
    if not changes:
        return jsonify(record=decorate(current, auth.store().users(), reasons_by_key()))
    result = auth.store().update_record(rid, {
        **changes, "updated_at": iso_now(auth.now()), "updated_by": g.user["user_id"]})
    if not result:
        raise ApiError("Record not found.", 404)
    before, after = result
    auth.audit("record_updated", target=rid, details={
        "owner": owner, "before": {k: before.get(k) for k in changes}, "after": changes})
    return jsonify(record=decorate(after, auth.store().users(), reasons_by_key()))


@bp.delete("/records/<rid>")
@auth.login_required
def delete_record(rid):
    _, owner = _owned_record_ref(rid)
    removed = auth.store().delete_record(rid)
    if not removed:
        raise ApiError("Record not found.", 404)
    auth.audit("record_deleted", target=rid, details={
        "owner": owner, **{k: removed.get(k) for k in
                           ("date", "time", "ip_phone", "reason", "note")}})
    return jsonify(ok=True)


# ===================================================================== admin
@bp.get("/admin/overview")
@auth.admin_required
def overview():
    d_from, d_to = parse_range()
    reasons = reasons_by_key()
    users = auth.store().users()
    recs = auth.store().records_between(d_from, d_to)
    rows = stats.per_employee(users, recs, reasons)
    totals = stats.summarize(recs, reasons)
    kind = "hour" if d_from == d_to else "day"
    recent = sorted(recs, key=lambda r: (r["date"], r["time"]), reverse=True)[:10]
    employees = [u for u in users.values() if u["role"] == "user" and u.get("active", True)]
    return jsonify(
        range={"from": d_from, "to": d_to}, server_time=iso_now(auth.now()),
        totals={k: totals[k] for k in ("total", *STATUSES)} | {
            "active_employees": sum(1 for r in rows if r["total"] > 0),
            "employees": len(employees)},
        chart={"kind": kind, "buckets": stats.buckets(recs, reasons, kind, d_from, d_to)},
        by_reason=totals["by_reason"], employees=rows,
        recent=[decorate(r, users, reasons) for r in recent],
        departments=sorted({u.get("department", "") for u in employees if u.get("department")}),
    )


@bp.get("/admin/records")
@auth.admin_required
def admin_records():
    d_from, d_to = parse_range()
    rows = filtered_records(d_from, d_to, request.args)
    for row in rows:
        row["editable"] = True
    return jsonify(**paginate(rows), range={"from": d_from, "to": d_to},
                   server_time=iso_now(auth.now()))


@bp.get("/admin/users")
@auth.admin_required
def list_users():
    users = sorted(auth.store().users().values(), key=lambda u: (u["role"] != "admin", u["name"]))
    return jsonify(users=[admin_user_row(u) for u in users])


def _clean_profile(body, partial=False):
    out = {}
    if not partial or "name" in body:
        name = str(body.get("name") or "").strip()
        if not 1 <= len(name) <= 80:
            raise ApiError("Enter the employee's name (up to 80 characters).")
        out["name"] = name
    if not partial or "department" in body:
        dept = str(body.get("department") or "").strip()
        if len(dept) > 80:
            raise ApiError("Department is too long (max 80 characters).")
        out["department"] = dept
    if "role" in body:
        if body["role"] not in ROLES:
            raise ApiError("Role must be 'user' or 'admin'.")
        out["role"] = body["role"]
    return out


def _new_user(uid, profile, password):
    return {"user_id": uid, "role": "user", "active": True, "must_change_password": True,
            "created_at": iso_now(auth.now()), "last_login": None,
            "password_hash": hash_password(password), **profile}


@bp.post("/admin/users")
@auth.admin_required
def create_user():
    body = json_body()
    uid = norm_uid(body.get("user_id"))
    if not valid_uid(uid):
        raise ApiError("User ID may use letters, digits, dot, dash and underscore (max 32).")
    profile = _clean_profile(body)
    password = str(body.get("password") or "")
    if password:
        problem = password_problem(password)
        if problem:
            raise ApiError(problem)
    else:
        password = generate_password()
    try:
        user = auth.store().create_user(_new_user(uid, {"role": "user", **profile}, password))
    except StorageError as exc:
        raise ApiError(str(exc), 409)
    auth.audit("user_created", target=uid, details={"name": user["name"], "role": user["role"]})
    return jsonify(user=admin_user_row(user), temp_password=password), 201


@bp.put("/admin/users/<uid>")
@auth.admin_required
def update_user(uid):
    uid = norm_uid(uid)
    body = json_body()
    changes = _clean_profile(body, partial=True)
    if "active" in body:
        changes["active"] = bool(body["active"])

    def guard(users, after):
        if uid == g.user["user_id"] and (not after.get("active", True) or after["role"] != "admin"):
            raise ApiError("You cannot disable or demote your own account.", 409)
        if not any(u.get("active", True) and u["role"] == "admin"
                   for k, u in users.items() if k != uid) and not (
                after.get("active", True) and after["role"] == "admin"):
            raise ApiError("At least one active administrator is required.", 409)

    result = auth.store().update_user(uid, changes, guard)
    if not result:
        raise ApiError("User not found.", 404)
    before, after = result
    diff = {k: {"from": before.get(k), "to": after.get(k)} for k in changes
            if before.get(k) != after.get(k)}
    if diff:
        auth.audit("user_updated", target=uid, details=diff)
    return jsonify(user=admin_user_row(after))


@bp.post("/admin/users/<uid>/reset-password")
@auth.admin_required
def reset_password(uid):
    uid = norm_uid(uid)
    body = request.get_json(silent=True) or {}
    password = str(body.get("password") or "")
    if password:
        problem = password_problem(password)
        if problem:
            raise ApiError(problem)
    else:
        password = generate_password()
    result = auth.store().update_user(uid, {
        "password_hash": hash_password(password), "must_change_password": True})
    if not result:
        raise ApiError("User not found.", 404)
    auth.audit("password_reset", target=uid)
    return jsonify(temp_password=password)


# ------------------------------------------------------------- bulk import
_HEADER_ALIASES = {
    "user_id": {"user_id", "userid", "user id", "id", "employee id", "employee_id", "username"},
    "name": {"name", "full name", "employee name", "employee"},
    "department": {"department", "dept", "section"},
    "password": {"password", "pwd"},
}


def _read_rows(upload):
    """Rows of a .csv or .xlsx upload as lists of strings."""
    filename = (upload.filename or "").lower()
    raw = upload.read()
    if filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        return [["" if c is None else str(c).strip() for c in row]
                for row in wb.worksheets[0].iter_rows(values_only=True)]
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise ApiError("Upload a .csv or .xlsx file.")


@bp.post("/admin/users/import")
@auth.admin_required
def import_users():
    upload = request.files.get("file")
    if not upload:
        raise ApiError("Choose a file to import.")
    rows = [r for r in _read_rows(upload) if any(r)]
    if not rows:
        raise ApiError("The file is empty.")
    header = [h.strip().lower() for h in rows[0]]
    col = {}
    for field, aliases in _HEADER_ALIASES.items():
        for i, h in enumerate(header):
            if h in aliases:
                col[field] = i
                break
    if "user_id" not in col or "name" not in col:
        raise ApiError("The first row must have column titles including 'user_id' and 'name' "
                       "(optional: 'department', 'password').")

    def cell(row, field):
        i = col.get(field)
        return row[i].strip() if i is not None and i < len(row) else ""

    to_create, skipped, seen, credentials = [], [], set(), []
    for line_no, row in enumerate(rows[1:], start=2):
        uid = norm_uid(cell(row, "user_id").removesuffix(".0"))  # Excel numbers arrive as 230966.0
        name, dept, password = cell(row, "name"), cell(row, "department"), cell(row, "password")
        problem = None
        if not valid_uid(uid):
            problem = "Invalid User ID"
        elif not name or len(name) > 80 or len(dept) > 80:
            problem = "Missing or too-long name/department"
        elif uid in seen:
            problem = "Duplicate User ID in the file"
        elif password and password_problem(password):
            problem = password_problem(password)
        if problem:
            skipped.append({"row": line_no, "user_id": uid, "reason": problem})
            continue
        seen.add(uid)
        password = password or generate_password()
        to_create.append(_new_user(uid, {"name": name, "department": dept}, password))
        credentials.append({"user_id": uid, "name": name, "password": password})

    created, existing = auth.store().create_users(to_create)
    created_ids = {u["user_id"] for u in created}
    for uid in existing:
        skipped.append({"row": None, "user_id": uid, "reason": "User ID already exists"})
    credentials = [c for c in credentials if c["user_id"] in created_ids]
    auth.audit("users_imported", details={"created": len(created), "skipped": len(skipped)})
    return jsonify(created=len(created), skipped=skipped, credentials=credentials)


# ======================================================================== audit
@bp.get("/admin/audit")
@auth.admin_required
def audit_log():
    d_from, d_to = parse_range()
    actor = norm_uid(request.args.get("actor"))
    action = request.args.get("action") or ""
    q = (request.args.get("q") or "").strip().lower()
    entries = auth.store().read_audit(d_from, d_to)
    actions = sorted(Counter(e.get("action") for e in entries))
    out = []
    for e in entries:
        if actor and actor not in (e.get("actor") or "").lower() \
                and actor not in (e.get("actor_name") or "").lower():
            continue
        if action and e.get("action") != action:
            continue
        if q and q not in re.sub(r"\s+", " ", str(e)).lower():
            continue
        out.append(e)
    out.reverse()  # files are chronological; show newest first
    return jsonify(**paginate(out, default_size=100), actions=actions)
