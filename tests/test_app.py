import io
import threading

import openpyxl
import pytest

from calllog.storage import Store
from conftest import Browser


# ------------------------------------------------------------------- login
def test_first_run_creates_admin_who_must_change_password(app):
    pw = app.config["INITIAL_ADMIN_PASSWORD"]
    assert pw and len(pw) >= 10
    b = Browser(app)
    resp = b.login("admin", pw)
    assert resp.status_code == 302
    assert b.get("/api/me").status_code == 403           # blocked until the password changes
    assert b.client.get("/admin").headers["Location"].endswith("/account/password")
    resp = b.client.post("/account/password", data={
        "current_password": pw, "new_password": "short", "confirm_password": "short",
        "csrf_token": b.csrf})
    assert resp.status_code == 400
    resp = b.client.post("/account/password", data={
        "current_password": pw, "new_password": "BrandNewPass1", "confirm_password": "BrandNewPass1",
        "csrf_token": b.csrf})
    assert resp.status_code == 302
    assert b.get("/api/me").get_json()["user"]["role"] == "admin"
    assert app.test_client().post("/login", data={"user_id": "admin", "password": pw}).status_code == 400


def test_wrong_password_is_rejected_and_throttled(app):
    b = Browser(app)
    for _ in range(5):
        assert b.login("admin", "nope").status_code == 401
    assert b.login("admin", app.config["INITIAL_ADMIN_PASSWORD"]).status_code == 429


def test_login_needs_csrf_token(app):
    resp = app.test_client().post("/login", data={"user_id": "admin", "password": "x"})
    assert resp.status_code == 400


def test_api_post_without_csrf_header_is_rejected(admin):
    resp = admin.client.post("/api/admin/users", json={"user_id": "1", "name": "x"})
    assert resp.status_code == 400


def test_anonymous_gets_401_and_redirect(app):
    c = app.test_client()
    assert c.get("/api/me").status_code == 401
    assert c.get("/admin").headers["Location"].endswith("/login")


# ---------------------------------------------------------------- records
def log_call(b, ip="75942", reason="solved", note=""):
    return b.post("/api/records", {"ip_phone": ip, "reason": reason, "note": note})


def test_employee_logs_edits_and_deletes_calls(make_employee, clock):
    emp = make_employee("230966", "Khaled")
    assert log_call(emp, "75942", "solved").status_code == 201
    clock.advance(minutes=30)
    assert log_call(emp, "٩٥٠٧٧", "answered").status_code == 201      # Arabic-Indic digits
    clock.advance(hours=1)
    rid = log_call(emp, "75071", "followup", "call back tomorrow").get_json()["record"]["id"]

    day = emp.get("/api/records").get_json()
    assert [r["ip_phone"] for r in day["records"]] == ["75071", "95077", "75942"]   # newest first
    s = day["stats"]
    assert (s["total"], s["resolved"], s["transferred"], s["ticket_inquiry"]) == (3, 2, 0, 1)
    assert s["first_call"] == "2026-10-11 09:00:00" and s["last_call"] == "2026-10-11 10:30:00"
    assert {h["key"]: h["total"] for h in s["hourly"] if h["total"]} == {9: 2, 10: 1}

    r = emp.put(f"/api/records/{rid}", {"reason": "transferred", "note": "sent to network team"})
    assert r.status_code == 200 and r.get_json()["record"]["status"] == "transferred"
    assert r.get_json()["record"]["edited"] is True
    assert emp.get("/api/records").get_json()["stats"]["transferred"] == 1

    assert emp.delete(f"/api/records/{rid}").status_code == 200
    assert emp.get("/api/records").get_json()["stats"]["total"] == 2
    assert emp.delete(f"/api/records/{rid}").status_code == 404


@pytest.mark.parametrize("payload", [
    {"ip_phone": "", "reason": "solved"},
    {"ip_phone": "abc", "reason": "solved"},
    {"ip_phone": "12", "reason": "solved"},
    {"ip_phone": "75942", "reason": "nonsense"},
    {"ip_phone": "75942", "reason": "solved", "note": "x" * 301},
])
def test_invalid_calls_are_rejected(make_employee, payload):
    assert make_employee("1001").post("/api/records", payload).status_code == 400


def test_employee_cannot_touch_other_employees_records(make_employee):
    a, b = make_employee("1001"), make_employee("1002")
    rid = log_call(a).get_json()["record"]["id"]
    assert b.put(f"/api/records/{rid}", {"note": "hacked"}).status_code == 404
    assert b.delete(f"/api/records/{rid}").status_code == 404
    assert a.get("/api/records").get_json()["records"][0]["note"] == ""
    assert b.get("/api/records").get_json()["stats"]["total"] == 0


@pytest.mark.parametrize("rid", ["../../users", "2026-10-11_../../x_abc", "garbage", "2026-13-40_1001_abc"])
def test_malformed_record_ids_are_404(make_employee, rid):
    emp = make_employee("1001")
    assert emp.delete(f"/api/records/{rid}").status_code == 404
    assert emp.put(f"/api/records/{rid}", {"note": "x"}).status_code == 404


def test_edit_window_locks_old_records_for_employees_not_admins(make_app, clock):
    app = make_app(user_edit_days=1)
    adm = Browser(app)
    adm.login("admin", app.config["INITIAL_ADMIN_PASSWORD"])
    adm.client.post("/account/password", data={"current_password": app.config["INITIAL_ADMIN_PASSWORD"],
                    "new_password": "AdminPass123", "confirm_password": "AdminPass123", "csrf_token": adm.csrf})
    adm.post("/api/admin/users", {"user_id": "1001", "name": "E", "password": "TempPass123"})
    emp = Browser(app)
    emp.login("1001", "TempPass123")
    emp.client.post("/account/password", data={"current_password": "TempPass123", "new_password": "MyOwnPass123",
                    "confirm_password": "MyOwnPass123", "csrf_token": emp.csrf})
    rid = log_call(emp).get_json()["record"]["id"]
    clock.advance(days=2)
    assert emp.put(f"/api/records/{rid}", {"note": "late"}).status_code == 403
    assert emp.delete(f"/api/records/{rid}").status_code == 403
    assert adm.put(f"/api/records/{rid}", {"note": "fixed by admin"}).status_code == 200
    day = emp.get("/api/records?date=2026-10-11").get_json()
    assert day["records"][0]["editable"] is False and day["records"][0]["note"] == "fixed by admin"


def test_cannot_view_future_dates(make_employee):
    assert make_employee("1001").get("/api/records?date=2026-10-12").status_code == 400


def test_admin_account_cannot_log_calls(admin):
    assert log_call(admin).status_code == 403


# -------------------------------------------------------------- admin views
def test_admin_overview_records_and_filters(admin, make_employee, clock):
    a, b = make_employee("1001", "Aya", "Identity"), make_employee("1002", "Omar", "Billing")
    make_employee("1003", "Idle", "Billing")
    log_call(a, "11111", "solved")
    clock.advance(minutes=1)
    log_call(a, "22222", "transferred")
    clock.advance(hours=2)
    log_call(b, "33333", "followup", "ticket 55")

    ov = admin.get("/api/admin/overview").get_json()
    assert ov["totals"] == {"total": 3, "resolved": 1, "transferred": 1, "ticket_inquiry": 1,
                            "active_employees": 2, "employees": 3}
    assert ov["chart"]["kind"] == "hour" and len(ov["chart"]["buckets"]) == 24
    assert [e["user_id"] for e in ov["employees"]][:2] == ["1001", "1002"]
    assert {e["user_id"]: e["total"] for e in ov["employees"]}["1003"] == 0   # idle staff are listed
    assert ov["recent"][0]["ip_phone"] == "33333" and ov["recent"][0]["name"] == "Omar"
    assert ov["departments"] == ["Billing", "Identity"]

    q = lambda s: admin.get("/api/admin/records?" + s).get_json()
    assert q("")["total"] == 3
    assert [r["ip_phone"] for r in q("user_id=1001")["items"]] == ["22222", "11111"]
    assert q("reason=followup")["total"] == 1
    assert q("q=ticket")["items"][0]["user_id"] == "1002"
    assert q("department=Billing")["total"] == 1
    assert q("page_size=2")["pages"] == 2

    week = admin.get("/api/admin/overview?from=2026-10-05&to=2026-10-11").get_json()
    assert week["chart"]["kind"] == "day" and len(week["chart"]["buckets"]) == 7


def test_admin_can_edit_and_delete_any_record(admin, make_employee):
    emp = make_employee("1001")
    rid = log_call(emp).get_json()["record"]["id"]
    assert admin.put(f"/api/records/{rid}", {"reason": "answered"}).status_code == 200
    assert admin.delete(f"/api/records/{rid}").status_code == 200


def test_employee_is_forbidden_from_admin_endpoints(make_employee):
    emp = make_employee("1001")
    for url in ("/api/admin/overview", "/api/admin/records", "/api/admin/users",
                "/api/admin/audit", "/export/records.csv", "/export/report.xlsx"):
        assert emp.get(url).status_code == 403, url
    assert emp.client.get("/admin").headers["Location"].endswith("/")


def test_bad_date_ranges(admin):
    assert admin.get("/api/admin/overview?from=2026-10-12&to=2026-10-01").status_code == 400
    assert admin.get("/api/admin/overview?from=nope").status_code == 400
    assert admin.get("/api/admin/overview?from=2020-01-01&to=2026-01-01").status_code == 400


# ---------------------------------------------------------- user management
def test_user_management_flow(admin, app):
    r = admin.post("/api/admin/users", {"user_id": "230966", "name": "خالد محمد", "department": "إدارة اعمال الهوية"})
    assert r.status_code == 201
    temp = r.get_json()["temp_password"]
    assert len(temp) >= 10
    assert admin.post("/api/admin/users", {"user_id": "230966", "name": "dup"}).status_code == 409
    assert admin.post("/api/admin/users", {"user_id": "bad id!", "name": "x"}).status_code == 400
    assert admin.post("/api/admin/users", {"user_id": "5", "name": ""}).status_code == 400

    emp = Browser(app)
    assert emp.login("230966", temp).status_code == 302
    assert emp.get("/api/me").status_code == 403       # must change the temporary password

    assert admin.put("/api/admin/users/230966", {"active": False}).status_code == 200
    assert Browser(app).login("230966", temp).status_code == 403
    assert emp.get("/api/records").status_code == 401   # an open session dies when disabled
    admin.put("/api/admin/users/230966", {"active": True, "department": "Billing"})

    new_pw = admin.post("/api/admin/users/230966/reset-password", {}).get_json()["temp_password"]
    assert Browser(app).login("230966", temp).status_code == 401
    assert Browser(app).login("230966", new_pw).status_code == 302
    users = {u["user_id"]: u for u in admin.get("/api/admin/users").get_json()["users"]}
    assert users["230966"]["department"] == "Billing" and "password_hash" not in users["230966"]


def test_admin_cannot_lock_themselves_out(admin):
    assert admin.put("/api/admin/users/admin", {"active": False}).status_code == 409
    assert admin.put("/api/admin/users/admin", {"role": "user"}).status_code == 409
    assert admin.put("/api/admin/users/ghost", {"name": "x"}).status_code == 404


def test_import_users_from_csv(admin, app):
    data = "user_id,name,department,password\n230966,خالد,Identity,\n230967,Sara,Billing,Passw0rdX\n,NoId,x,\n230966,Dup,x,\n"
    r = admin.post("/api/admin/users/import", data={"file": (io.BytesIO(data.encode("utf-8")), "u.csv")},
                   content_type="multipart/form-data")
    body = r.get_json()
    assert r.status_code == 200 and body["created"] == 2 and len(body["skipped"]) == 2
    creds = {c["user_id"]: c["password"] for c in body["credentials"]}
    assert creds["230967"] == "Passw0rdX" and len(creds["230966"]) >= 10
    assert Browser(app).login("230966", creds["230966"]).status_code == 302
    assert admin.post("/api/admin/users/import", data={"file": (io.BytesIO(b"x"), "u.txt")},
                      content_type="multipart/form-data").status_code == 400


# ------------------------------------------------------------------- audit
def test_audit_log_records_who_did_what(admin, make_employee):
    emp = make_employee("1001", "Aya")
    rid = log_call(emp, "75942", "solved").get_json()["record"]["id"]
    emp.put(f"/api/records/{rid}", {"ip_phone": "75943", "note": "typo"})
    emp.delete(f"/api/records/{rid}")
    Browser(admin.client.application).login("1001", "wrong-password")

    entries = admin.get("/api/admin/audit").get_json()
    actions = [e["action"] for e in entries["items"]]
    for expected in ("login", "record_created", "record_updated", "record_deleted",
                     "login_failed", "user_created", "password_changed"):
        assert expected in actions, expected
    upd = next(e for e in entries["items"] if e["action"] == "record_updated")
    assert upd["actor"] == "1001" and upd["details"]["before"]["ip_phone"] == "75942"
    assert upd["details"]["after"] == {"ip_phone": "75943", "note": "typo"}
    deleted = next(e for e in entries["items"] if e["action"] == "record_deleted")
    assert deleted["details"]["ip_phone"] == "75943"                  # the deleted data is kept
    only = admin.get("/api/admin/audit?action=record_created&actor=aya").get_json()
    assert only["total"] == 1


# ----------------------------------------------------------------- exports
def test_exports(admin, make_employee):
    emp = make_employee("1001", "Aya", "Identity")
    log_call(emp, "75942", "solved", "=HYPERLINK(\"http://evil\")")
    log_call(emp, "75943", "followup")

    r = admin.get("/export/report.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Calls"]
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert summary[0][:4] == ("User ID", "Employee", "Department", "Total calls")
    assert summary[1][:7] == ("1001", "Aya", "Identity", 2, 1, 0, 1)
    rows = list(wb["Calls"].iter_rows(values_only=True))
    assert rows[0][5:9] == ("IP Phone", "Action Taken", "Action (English)", "Status")
    by_ip = {r[5]: r for r in rows[1:]}
    assert by_ip["75942"][6] == "تم حل المشكلة" and by_ip["75942"][8] == "Resolved"
    assert by_ip["75943"][8] == "Ticket Inquiry"
    formula_cells = [c for row in wb["Calls"].iter_rows() for c in row if c.data_type == "f"]
    assert by_ip["75942"][9].startswith("=HYPERLINK") and not formula_cells    # stored as text, never a live formula

    csv_resp = admin.get("/export/records.csv")
    text = csv_resp.data.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Date,Time,User ID") and "'=HYPERLINK" in text
    assert "تم حل المشكلة" in text
    assert admin.get("/export/audit.xlsx").status_code == 200
    assert admin.get("/export/audit.csv").status_code == 200
    assert admin.get("/export/records.pdf").status_code == 404

    mine = emp.get("/export/my.xlsx")
    assert mine.status_code == 200
    assert len(list(openpyxl.load_workbook(io.BytesIO(mine.data)).active.iter_rows())) == 3
    actions = [e["action"] for e in admin.get("/api/admin/audit").get_json()["items"]]
    assert actions.count("export") >= 4


# ----------------------------------------------------------------- storage
def test_concurrent_writes_do_not_lose_records(tmp_path):
    store = Store(tmp_path / "d")
    store.init()

    def work(n):
        for i in range(25):
            store.add_record({"id": f"2026-10-11_u{n}_{i:06x}", "user_id": f"u{n}", "date": "2026-10-11",
                              "time": "09:00:00", "ip_phone": "12345", "reason": "solved", "note": ""})
            # also hammer one shared file
            store.add_record({"id": f"2026-10-11_shared_{n}{i:05x}", "user_id": "shared",
                              "date": "2026-10-11", "time": "09:00:01", "ip_phone": "1", "reason": "solved", "note": ""})

    threads = [threading.Thread(target=work, args=(n,)) for n in range(8)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert len(store.records_between("2026-10-11", "2026-10-11")) == 8 * 25 * 2
    assert len(store.day_records("2026-10-11", "shared")) == 200


def test_unreachable_data_folder_gives_a_clear_error(tmp_path):
    from calllog.storage import StorageError
    blocker = tmp_path / "file"
    blocker.write_text("x")
    with pytest.raises(StorageError, match="Cannot use the data folder"):
        Store(blocker / "sub").init()


def test_security_headers(app):
    r = app.test_client().get("/login")
    assert "script-src 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["X-Frame-Options"] == "DENY"
